import json
import sqlite3
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from retail_insights.pipeline import (
    aggregate, build_database, clean_orders, enrich_orders, normalize_references,
    parse_money, reconcile_sql, run_analysis, summarize,
)


def example_tables():
    customers = pd.DataFrame([
        {"customer_id": "C1", "city": "Dhaka", "region": "Dhaka", "segment": "Consumer"},
        {"customer_id": "C2", "city": "Khulna", "region": "Khulna", "segment": "Small Business"},
    ])
    products = pd.DataFrame([
        {"product_id": "P1", "product_name": "Mouse", "category": "Electronics", "unit_cost_bdt": "6.00"},
        {"product_id": "P2", "product_name": "Book", "category": "Books", "unit_cost_bdt": "20.00"},
    ])
    rows = [
        ["A1", "2026-01-01", "C1", "P1", "3", "10.01", "15", "completed"],
        ["A2", "2026-01-02", "C2", "P2", "2", "50.00", "10", "completed"],
        ["A3", "2026-02-01", "C2", "P2", "2", "50.00", "0", "returned"],
        ["A4", "2026-02-02", "C1", "P1", "1", "10.00", "0", "cancelled"],
        ["A5", "2026-02-10", "C1", "P1", "1", "10.00", "100", "completed"],
    ]
    raw = pd.DataFrame(rows, columns=["order_id", "order_date", "customer_id", "product_id",
                                     "quantity", "unit_price_bdt", "discount_pct", "status"])
    return customers, products, raw


class CleaningTests(unittest.TestCase):
    def setUp(self):
        customers, products, self.raw = example_tables()
        self.customers, self.products = normalize_references(customers, products)

    def test_currency_formats_and_invalid_precision(self):
        self.assertEqual(parse_money(" BDT 1,250.50 "), 125050)
        self.assertEqual(parse_money("৳ 12.05"), 1205)
        for value in ("NaN", "Infinity", "-1", "0.001", "", "unknown"):
            with self.subTest(value=value):
                self.assertIsNone(parse_money(value))

    def test_normalizes_whitespace_case_and_preserves_ids(self):
        self.raw.loc[0, ["order_id", "customer_id", "product_id", "status"]] = [" a1 ", " c1 ", " p1 ", " COMPLETED "]
        clean, rejected = clean_orders(self.raw, self.customers, self.products)
        self.assertTrue(rejected.empty)
        self.assertEqual(clean.iloc[0]["order_id"], "A1")
        self.assertEqual(clean.iloc[0]["unit_price_paisa"], 1001)

    def test_duplicate_ids_are_audited_without_double_counting(self):
        raw = pd.concat([self.raw, self.raw.iloc[[0]]], ignore_index=True)
        clean, rejected = clean_orders(raw, self.customers, self.products)
        self.assertEqual(len(clean), 5)
        self.assertEqual(rejected.iloc[0]["reason"], "duplicate_order_id")
        self.assertEqual(rejected.iloc[0]["source_row"], 7)

    def test_invalid_row_does_not_hide_a_later_valid_id(self):
        first = self.raw.iloc[[0]].copy()
        first.loc[0, "quantity"] = "invalid"
        raw = pd.concat([first, self.raw], ignore_index=True)
        clean, rejected = clean_orders(raw, self.customers, self.products)
        self.assertEqual(len(clean), 5)
        self.assertEqual(len(rejected), 1)

    def test_rejects_invalid_dates_and_values_with_reasons(self):
        cases = [("order_date", "2026-02-30", "invalid_date"),
                 ("order_date", "20260101", "invalid_date"),
                 ("quantity", "1.5", "invalid_quantity"), ("quantity", "0", "invalid_quantity"),
                 ("discount_pct", "101", "invalid_discount"),
                 ("unit_price_bdt", "-1", "invalid_price"),
                 ("customer_id", "C999", "unknown_customer"),
                 ("product_id", "P999", "unknown_product"),
                 ("status", "pending", "invalid_status"),
                 ("order_id", "", "missing_order_id")]
        for column, value, reason in cases:
            with self.subTest(column=column, value=value):
                raw = self.raw.iloc[[0]].copy()
                raw.loc[0, column] = value
                clean, rejected = clean_orders(raw, self.customers, self.products)
                self.assertTrue(clean.empty)
                self.assertIn(reason, rejected.iloc[0]["reason"])

    def test_missing_columns_fail_clearly(self):
        with self.assertRaisesRegex(ValueError, "missing columns: quantity"):
            clean_orders(self.raw.drop(columns="quantity"), self.customers, self.products)

    def test_duplicate_reference_ids_fail_before_join(self):
        customers, products, _ = example_tables()
        customers = pd.concat([customers, customers.iloc[[0]]])
        with self.assertRaisesRegex(ValueError, "Duplicate customer_id"):
            normalize_references(customers, products)

    def test_invalid_reference_cost_is_rejected(self):
        customers, products, _ = example_tables()
        products.loc[0, "unit_cost_bdt"] = "-5"
        with self.assertRaisesRegex(ValueError, "unit costs"):
            normalize_references(customers, products)


class AnalysisTests(unittest.TestCase):
    def setUp(self):
        customers, products, raw = example_tables()
        self.customers, self.products = normalize_references(customers, products)
        self.clean, _ = clean_orders(raw, self.customers, self.products)
        self.data = enrich_orders(self.clean, self.customers, self.products)

    def test_half_up_rounding_after_discount(self):
        order = self.data.set_index("order_id").loc["A1"]
        self.assertEqual(order["net_sales_paisa"], 2553)
        self.assertEqual(order["gross_profit_paisa"], 753)

    def test_returns_and_cancellations_book_no_sales_or_cost(self):
        excluded = self.data[self.data["status"].ne("completed")]
        self.assertEqual(excluded["net_sales_paisa"].sum(), 0)
        self.assertEqual(excluded["cost_paisa"].sum(), 0)

    def test_full_discount_retains_cost_and_negative_profit(self):
        order = self.data.set_index("order_id").loc["A5"]
        self.assertEqual(order["net_sales_paisa"], 0)
        self.assertEqual(order["gross_profit_paisa"], -600)

    def test_kpis_match_hand_calculated_values(self):
        result = summarize(self.data, 7, 2)
        self.assertEqual(result["net_sales_bdt"], 115.53)
        self.assertEqual(result["gross_profit_bdt"], 51.53)
        self.assertEqual(result["average_order_value_bdt"], 38.51)
        self.assertEqual(result["return_rate_pct"], 25.0)
        self.assertEqual(result["completed_orders"], 3)
        self.assertEqual(result["units_sold"], 6)

    def test_sql_and_pandas_agree_for_all_summary_groups(self):
        with sqlite3.connect(":memory:") as connection:
            queries = build_database(connection, self.clean, self.customers, self.products)
            reconcile_sql(self.data, queries)
            self.assertEqual(int(queries["kpis"].iloc[0]["net_sales_paisa"]), 11553)
            self.assertEqual(len(queries), 8)
            self.assertEqual(queries["above_average_orders"]["order_id"].tolist(), ["A2"])

    def test_mismatched_sql_result_stops_reporting(self):
        with sqlite3.connect(":memory:") as connection:
            queries = build_database(connection, self.clean, self.customers, self.products)
        queries["kpis"].loc[0, "net_sales_paisa"] += 1
        with self.assertRaisesRegex(ValueError, "totals disagree"):
            reconcile_sql(self.data, queries)

    def test_database_enforces_foreign_keys(self):
        with sqlite3.connect(":memory:") as connection:
            build_database(connection, self.clean, self.customers, self.products)
            with self.assertRaises(sqlite3.IntegrityError):
                connection.execute("UPDATE sales SET customer_id = 'UNKNOWN' WHERE order_id = 'A1'")

    def test_no_completed_orders_has_defined_zero_totals(self):
        clean = self.clean.copy()
        clean["status"] = "returned"
        data = enrich_orders(clean, self.customers, self.products)
        with sqlite3.connect(":memory:") as connection:
            queries = build_database(connection, clean, self.customers, self.products)
            reconcile_sql(data, queries)
        summary = summarize(data, 5, 0)
        self.assertEqual(summary["net_sales_bdt"], 0)
        self.assertIsNone(summary["gross_margin_pct"])
        self.assertEqual(summary["return_rate_pct"], 100)


class ReportTests(unittest.TestCase):
    def test_full_pipeline_and_repeat_run(self):
        customers, products, raw = example_tables()
        raw = pd.concat([raw, raw.iloc[[0]]], ignore_index=True)
        with tempfile.TemporaryDirectory() as folder:
            source, destination = Path(folder) / "raw", Path(folder) / "reports"
            source.mkdir()
            customers.to_csv(source / "customers.csv", index=False)
            products.to_csv(source / "products.csv", index=False)
            raw.to_csv(source / "sales_raw.csv", index=False)
            first = run_analysis(source, destination)
            second = run_analysis(source, destination)
            self.assertEqual(first, second)
            self.assertEqual(json.loads((destination / "kpis.json").read_text())["rejected_rows"], 1)
            with sqlite3.connect(destination / "analysis.sqlite") as connection:
                self.assertEqual(connection.execute("SELECT COUNT(*) FROM sales").fetchone()[0], 5)
            workbook = load_workbook(destination / "sales_analysis.xlsx")
            self.assertEqual(workbook["Clean_Sales"].max_row, 6)
            self.assertEqual(workbook["Formula_Practice"]["B2"].value, "=SUM(SalesData[net_sales_bdt])")
            self.assertEqual(len(workbook["Monthly"]._charts), 1)
            workbook.close()
            page = (destination / "dashboard.html").read_text()
            self.assertNotIn("__RECORDS_JSON__", page)
            self.assertIn('id="sales-data"', page)
            self.assertEqual(len(list((destination / "charts").glob("*.png"))), 3)
            self.assertEqual(len(list((destination / "sql_results").glob("*.csv"))), 8)
            self.assertEqual(len(pd.read_csv(destination / "bi" / "fact_sales.csv")), 5)

    def test_input_output_overlap_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "must not overlap"):
            run_analysis(Path("data/raw"), Path("data/raw/reports"))


if __name__ == "__main__":
    unittest.main()
