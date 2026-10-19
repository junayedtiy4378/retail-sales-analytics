"""Export an offline dashboard, charts, Excel workbook, and BI tables."""

import json
import os
from pathlib import Path

import pandas as pd

from retail_insights.pipeline import ROOT, aggregate, export_money


def write_dashboard(destination, data, summary):
    columns = ["order_id", "order_date", "month", "region", "category", "product_name",
               "quantity", "status", "net_sales_paisa", "gross_profit_paisa"]
    records = data[columns].to_json(orient="records", force_ascii=True)
    template = (Path(__file__).parent / "dashboard.html").read_text(encoding="utf-8")
    template = template.replace("__RECORDS_JSON__", records.replace("<", "\\u003c"))
    template = template.replace("__SUMMARY_JSON__", json.dumps(summary).replace("<", "\\u003c"))
    (destination / "dashboard.html").write_text(template, encoding="utf-8")


def write_charts(destination, data):
    os.environ.setdefault("MPLCONFIGDIR", str(ROOT / ".cache" / "matplotlib"))
    os.environ.setdefault("XDG_CACHE_HOME", str(ROOT / ".cache"))
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    charts = destination / "charts"
    charts.mkdir(exist_ok=True)
    definitions = [
        ("month", "net_sales_paisa", "Monthly net sales", "monthly_sales.png", "#087f8c"),
        ("region", "net_sales_paisa", "Net sales by region", "regional_sales.png", "#3156a6"),
        ("category", "gross_profit_paisa", "Gross profit by category", "category_profit.png", "#7a5195"),
    ]
    for key, measure, title, filename, color in definitions:
        table = aggregate(data, key)
        fig, axis = plt.subplots(figsize=(9, 4.8), layout="constrained")
        if table.empty:
            axis.text(0.5, 0.5, "No completed orders", ha="center", transform=axis.transAxes)
        else:
            axis.bar(table[key], table[measure] / 100, color=color, width=0.6)
        axis.set_title(title + " — synthetic sample", loc="left", pad=16, fontweight="bold")
        axis.set_ylabel("BDT")
        axis.spines[["top", "right"]].set_visible(False)
        axis.grid(axis="y", alpha=0.2)
        axis.set_axisbelow(True)
        axis.ticklabel_format(axis="y", style="plain", useOffset=False)
        fig.savefig(charts / filename, dpi=160)
        plt.close(fig)


def write_excel(destination, data, customers, products, rejected, summary, queries):
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.workbook.properties import CalcProperties

    notes = [
        ("Dataset", "Original synthetic retail sample; not a real business."),
        ("Grain", "One row per order, with one product per order."),
        ("Currency", "BDT. Source calculations use integer paisa; 100 paisa = 1 BDT."),
        ("Sales", "Completed orders only, after discount; half-up rounding per order."),
        ("Profit", "Net sales minus product cost; excludes delivery, tax, and fees."),
        ("Returns", "Returned / (completed + returned). Cancelled orders excluded."),
        ("Formulas", "Python writes formulas; Excel recalculates them when opened."),
        ("Pivot practice", "Use Clean_Sales to insert your own Excel PivotTable."),
    ]
    tables = {
        "Read_Me": pd.DataFrame(notes, columns=["topic", "definition"]),
        "KPIs": pd.DataFrame(summary.items(), columns=["metric", "value"]),
        "Clean_Sales": export_money(data),
        "Monthly": export_money(queries["monthly_sales"]),
        "Regions": export_money(queries["regional_sales"]),
        "Categories": export_money(queries["category_sales"]),
        "Products": export_money(products),
        "Customers": customers,
        "Rejected_Rows": rejected,
    }
    with pd.ExcelWriter(destination / "sales_analysis.xlsx", engine="openpyxl") as writer:
        for sheet, frame in tables.items():
            frame.to_excel(writer, sheet_name=sheet, index=False)
            ws = writer.sheets[sheet]
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = PatternFill("solid", fgColor="16324F")
            if len(frame):
                table_name = "SalesData" if sheet == "Clean_Sales" else "Table" + sheet.replace("_", "")
                table = Table(displayName=table_name, ref=ws.dimensions)
                table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
                ws.add_table(table)
            for column in ws.columns:
                name = column[0].column_letter
                lengths = [len(str(cell.value or "")) for cell in list(column)[:60]]
                ws.column_dimensions[name].width = min(38, max(14, max(lengths, default=14) + 2))
                if str(column[0].value).endswith("_bdt"):
                    for cell in list(column)[1:]:
                        cell.number_format = '#,##0.00'
        formula = writer.book.create_sheet("Formula_Practice")
        for row in [
            ["Metric", "Formula / result"],
            ["Net sales (BDT)", "=SUM(SalesData[net_sales_bdt])"],
            ["Completed orders", '=COUNTIF(SalesData[status],"completed")'],
            ["Average order value (BDT)", "=IFERROR(B2/B3,0)"],
            ["Dhaka net sales (BDT)", '=SUMIFS(SalesData[net_sales_bdt],SalesData[region],"Dhaka")'],
            ["Gross profit (BDT)", "=SUM(SalesData[gross_profit_bdt])"],
            ["Gross margin", "=IFERROR(B6/B2,0)"],
        ]:
            formula.append(row)
        formula.column_dimensions["A"].width = 32
        formula.column_dimensions["B"].width = 60
        formula.freeze_panes = "A2"
        for row in (2, 4, 5, 6):
            formula[f"B{row}"].number_format = '#,##0.00'
        formula["B7"].number_format = "0.00%"
        writer.book.calculation = CalcProperties(fullCalcOnLoad=True)
        ws = writer.sheets["Monthly"]
        if len(queries["monthly_sales"]):
            chart = BarChart()
            chart.title = "Monthly net sales — synthetic sample"
            chart.y_axis.title = "BDT"
            chart.add_data(Reference(ws, min_col=3, min_row=1, max_row=ws.max_row), titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
            ws.add_chart(chart, "G2")


def write_insights(destination, summary, queries):
    monthly = queries["monthly_sales"]
    regional = queries["regional_sales"]
    categories = queries["category_sales"]
    lines = ["# Analysis findings", "", "**Synthetic demonstration data, not real business findings.**", "",
             f"- Raw rows: {summary['raw_rows']}; valid orders: {summary['clean_rows']}; rejected rows: {summary['rejected_rows']}.",
             f"- Completed orders: {summary['completed_orders']}; returned: {summary['returned_orders']}; cancelled: {summary['cancelled_orders']}.",
             f"- Net sales: BDT {summary['net_sales_bdt']:,.2f}.",
             f"- Gross profit before overheads: BDT {summary['gross_profit_bdt']:,.2f}.",
             f"- Return rate: {summary['return_rate_pct']:.2f}% (cancelled orders excluded)."]
    for label, key, table, measure in [
        ("Highest-sales month", "month", monthly, "net_sales_paisa"),
        ("Highest-sales region", "region", regional, "net_sales_paisa"),
        ("Highest-profit category", "category", categories, "gross_profit_paisa"),
    ]:
        if not table.empty:
            row = table.sort_values([measure, key], ascending=[False, True]).iloc[0]
            lines.append(f"- {label}: {row[key]} (BDT {int(row[measure]) / 100:,.2f}).")
    lines += ["", "## Interpretation limits", "",
              "The data was generated with a fixed random seed. Patterns illustrate analytical methods, not market demand or causal effects.",
              "Gross profit excludes shipping, payment fees, taxes, returns handling, and operating costs. Returned and cancelled orders have zero booked sales and cost in this simplified model.",
              "", "## Next analysis tasks", "",
              "1. Compare category margin percentages, not just revenue.",
              "2. Investigate return rates by region; check sample sizes first.",
              "3. Add a real, appropriately licensed dataset and document its definitions.",
              "4. Rebuild the charts in Power BI or Tableau using the BI guide."]
    (destination / "findings.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_reports(destination, data, customers, products, rejected, summary, queries):
    bi = destination / "bi"
    bi.mkdir(exist_ok=True)
    fact_columns = ["order_id", "order_date", "month", "customer_id", "product_id", "quantity",
                    "status", "net_sales_paisa", "cost_paisa", "gross_profit_paisa"]
    export_money(data[fact_columns]).to_csv(bi / "fact_sales.csv", index=False, float_format="%.2f")
    customers.to_csv(bi / "dim_customers.csv", index=False)
    export_money(products).to_csv(bi / "dim_products.csv", index=False, float_format="%.2f")
    write_dashboard(destination, data, summary)
    write_charts(destination, data)
    write_excel(destination, data, customers, products, rejected, summary, queries)
    write_insights(destination, summary, queries)
